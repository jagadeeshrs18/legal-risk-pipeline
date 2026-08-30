from __future__ import annotations

import json
import os
from datetime import datetime
from pathlib import Path
from uuid import uuid4

from docx import Document
from docx.shared import Inches
from flask import Flask, flash, redirect, render_template, request, send_file, url_for
from werkzeug.utils import secure_filename

from legalrisk.flask_analyzer import AcordRetriever, analyse, extract_pdf
from legalrisk.document_classifier import detect_document_category, detect_commercial_subtype

CATEGORIES = [
    "Business and Commercial Contracts",
    "Real Estate and Property Documents",
    "Estate Planning Documents",
    "Court and Litigation Pleadings",
]
COMMERCIAL_SUBTYPES = ["Employment Agreement", "NDA", "Service Agreement", "Vendor Agreement"]

# Perspective options shown per category/sub-type. For the four entries
# that also appear in legalrisk.flask_analyzer.PERSPECTIVE_FLIP, choosing
# the "opposite" side causes clauses written to protect the other side to
# be excluded from the risk score and shown as favourable instead.
PERSPECTIVE_OPTIONS = {
    "Employment Agreement": ["Employee", "Employer"],
    "NDA": ["Disclosing Party", "Receiving Party"],
    "Service Agreement": ["Customer", "Vendor"],
    "Vendor Agreement": ["Customer", "Vendor"],
    "Real Estate and Property Documents": ["Tenant", "Landlord"],
    "Estate Planning Documents": ["Beneficiary", "Executor", "Testator/Estate"],
    "Court and Litigation Pleadings": ["Plaintiff", "Defendant"],
}

ROOT = Path(__file__).resolve().parent
UPLOADS, REPORTS, CACHE = ROOT / "uploads", ROOT / "reports", ROOT / "indexes"
for folder in (UPLOADS, REPORTS, CACHE): folder.mkdir(exist_ok=True)

app = Flask(__name__)
app.secret_key = os.environ.get("FLASK_SECRET_KEY", "change-this-development-key")
app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024
_retriever = None


def get_retriever():
    global _retriever
    if _retriever is None:
        data_dir = Path(os.environ.get("ACORD_DATA_DIR", ROOT / "data" / "acord" / "extracted"))
        _retriever = AcordRetriever(data_dir, CACHE)
    return _retriever


@app.get("/")
def index():
    commercial_perspective_map = {st: PERSPECTIVE_OPTIONS[st] for st in COMMERCIAL_SUBTYPES}
    return render_template(
        "index.html",
        categories=CATEGORIES,
        commercial_subtypes=COMMERCIAL_SUBTYPES,
        perspective_options=PERSPECTIVE_OPTIONS,
        commercial_perspective_map=commercial_perspective_map,
    )


@app.post("/analyse")
def analyse_upload():
    uploaded = request.files.get("document")
    if not uploaded or not uploaded.filename:
        flash("Select a PDF document.", "danger"); return redirect(url_for("index"))
    if Path(uploaded.filename).suffix.lower() != ".pdf":
        flash("Only PDF files are accepted.", "danger"); return redirect(url_for("index"))

    selected_category = request.form.get("category", "Business and Commercial Contracts")
    if selected_category not in CATEGORIES:
        flash("Unknown document category selected.", "danger"); return redirect(url_for("index"))

    is_commercial = selected_category == "Business and Commercial Contracts"
    if is_commercial:
        contract_type = request.form.get("subtype", "Service Agreement")
        if contract_type not in COMMERCIAL_SUBTYPES:
            flash("Unknown document sub-type selected.", "danger"); return redirect(url_for("index"))
    else:
        contract_type = selected_category

    token = uuid4().hex
    path = UPLOADS / f"{token}_{secure_filename(uploaded.filename)}"
    uploaded.save(path)
    try:
        text, pages = extract_pdf(path)

        # Category gate: does this document actually look like the category
        # section the user dropped it into? If it confidently looks like a
        # *different* top-level category, reject it before running the
        # expensive embedding-based analysis and tell the user where it belongs.
        category_check = detect_document_category(text)
        if not category_check.matches(selected_category):
            flash(
                f"This document does not look like a {selected_category}. "
                f"It reads more like a {category_check.predicted} "
                f"(confidence {category_check.confidence:.0%}). "
                f"Please upload it in the correct category section below.",
                "warning",
            )
            return redirect(url_for("index"))

        # Soft check within Business & Commercial: does the chosen sub-type
        # (Employment/NDA/Service/Vendor) look right? This is only a
        # heads-up, not a hard rejection, since the top-level category is
        # already confirmed.
        if is_commercial:
            subtype_check = detect_commercial_subtype(text)
            if not subtype_check.matches(contract_type):
                flash(
                    f"Note: this looks more like a {subtype_check.predicted} than a "
                    f"{contract_type}. Continuing analysis using {contract_type} as selected — "
                    f"re-upload with the correct sub-type if this was a mistake.",
                    "info",
                )

        report = analyse(text, pages, get_retriever(), request.form.get("perspective", "Customer"),
                         contract_type,
                         request.form.get("jurisdiction", "India"))
        report["report_id"] = token; report["filename"] = uploaded.filename
        report["category"] = selected_category
        report["created_at"] = datetime.now().strftime("%d-%m-%Y %H:%M")
        report["category_confidence"] = category_check.confidence
        (REPORTS / f"{token}.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        return redirect(url_for("view_report", report_id=token))
    except Exception as exc:
        flash(str(exc), "danger"); return redirect(url_for("index"))
    finally:
        path.unlink(missing_ok=True)


def load_report(report_id):
    if not report_id.isalnum(): return None
    path = REPORTS / f"{report_id}.json"
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else None


@app.get("/report/<report_id>")
def view_report(report_id):
    report = load_report(report_id)
    if not report: return "Report not found", 404
    return render_template("report.html", report=report)


@app.get("/report/<report_id>/word")
def word_report(report_id):
    r = load_report(report_id)
    if not r: return "Report not found", 404
    doc = Document(); section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.65)
    doc.add_heading("Preliminary Legal Document Risk Assessment", 0)
    doc.add_paragraph(f"Document: {r['filename']}\nGenerated: {r['created_at']}\nPerspective: {r['perspective']}\nJurisdiction: {r['jurisdiction']}")
    doc.add_heading(f"Overall risk: {r['level']} ({r['overall_score']}/100)", 1)
    doc.add_paragraph(f"Pages: {r['pages']} | Analysed clauses: {r['clause_count']}")
    doc.add_heading("Missing protections", 1)
    if not r["missing"]: doc.add_paragraph("No required provision was flagged as missing.")
    for x in r["missing"]: doc.add_paragraph(f"{x['type']}: {x['reason']} (+{x['score']})", style="List Bullet")
    doc.add_heading("Clause findings", 1)
    for x in r["findings"]:
        if x.get("favors_you"):
            doc.add_heading(f"Clause {x['number']} — {x['type']} — favours you (raw risk {x['raw_score']})", 2)
        else:
            doc.add_heading(f"Clause {x['number']} — {x['type']} — score {x['score']}", 2)
        doc.add_paragraph(x["text"])
        doc.add_paragraph(f"Nearest ACORD similarity: {x['confidence']:.3f}. Similarity supports retrieval, not a legal conclusion.")
        if x.get("favors_you"):
            doc.add_paragraph("This clause type is typically drafted to protect the other side of this "
                               "relationship — from the selected perspective, it looks favourable rather "
                               "than risky, so it is not counted in the overall score.")
        elif x["risky_hits"]:
            doc.add_paragraph("Risk indicators: " + "; ".join(x["risky_hits"]))
        for s in x.get("suggestions", []):
            doc.add_paragraph(f"Suggested rewording — {s['issue']}", style="Intense Quote")
            doc.add_paragraph(s["suggested_clause"])
            doc.add_paragraph(f"Why: {s['rationale']}")
    doc.add_heading("Important limitation", 1)
    doc.add_paragraph("This automated research report is not legal advice. ACORD provides retrieval relevance, not risk labels. Every finding and missing-clause decision must be reviewed by a qualified legal professional, particularly for Indian law.")
    out = REPORTS / f"Legal_Risk_Report_{report_id}.docx"; doc.save(out)
    return send_file(out, as_attachment=True, download_name=out.name)


@app.errorhandler(413)
def too_large(_):
    flash("PDF exceeds the 20 MB upload limit.", "danger"); return redirect(url_for("index"))


if __name__ == "__main__": app.run(debug=True, host="127.0.0.1", port=5000)

